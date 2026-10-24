"""Exportacion del reporte de rentabilidad a Excel y a PDF (RF06).

Toma el diccionario que arma ReporteServicio.rentabilidad(), asi que las
cifras y los permisos son siempre los mismos que ve el usuario en pantalla.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TIPO_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIPO_PDF = "application/pdf"

# Paleta corporativa del proyecto.
VERDE = "4CAF50"
GRIS = "F5F5F5"

ENCABEZADOS = (
    "Contrato",
    "Lote",
    "Estado",
    "Cabezas",
    "Ingreso total",
    "Costo estimado",
    "Utilidad proyectada",
    "Recaudado",
    "Saldo por cobrar",
)
CAMPOS_VALOR = (
    "ingreso_total",
    "costo_estimado",
    "utilidad_proyectada",
    "recaudado",
    "saldo_por_cobrar",
)
FORMATO_PESOS = '"$" #,##0.00'


def numero_contrato(id_contrato: int) -> str:
    """Numero de contrato con ceros a la izquierda: 0001."""
    return f"{id_contrato:04d}"


def moneda(valor: float) -> str:
    """Pesos colombianos con punto de miles y coma decimal: $ 1.234.567,50."""
    texto = f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"$ {texto}"


def nombre_archivo(extension: str, ahora: datetime | None = None) -> str:
    """Nombre del archivo descargado: reporte_rentabilidad_20261024.xlsx."""
    ahora = ahora or datetime.now()
    return f"reporte_rentabilidad_{ahora:%Y%m%d}.{extension}"


def _titulo(reporte: dict) -> str:
    return f"Reporte de rentabilidad (margen del {reporte['margen_utilidad']:.0%})"


class ExportacionServicio:
    """Convierte el reporte de rentabilidad en un archivo descargable."""

    def excel(self, reporte: dict, ahora: datetime | None = None) -> bytes:
        """Libro de Excel con una hoja; los totales son formulas SUM."""
        ahora = ahora or datetime.now()
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Rentabilidad"
        hoja["A1"] = _titulo(reporte)
        hoja["A1"].font = Font(bold=True, size=14, color=VERDE)
        hoja["A2"] = f"Generado el {ahora:%d/%m/%Y %H:%M}"

        fila_encabezado = 4
        borde = Side(style="thin", color="BDBDBD")
        for columna, texto in enumerate(ENCABEZADOS, start=1):
            celda = hoja.cell(row=fila_encabezado, column=columna, value=texto)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=VERDE)
            celda.alignment = Alignment(horizontal="center", wrap_text=True)
            celda.border = Border(bottom=borde)

        fila = fila_encabezado
        for c in reporte["contratos"]:
            fila += 1
            valores = [
                numero_contrato(c["id_contrato"]),
                c["nombre_lote"],
                c["estado"],
                c["total_cabezas"],
                *(c[campo] for campo in CAMPOS_VALOR),
            ]
            for columna, valor in enumerate(valores, start=1):
                celda = hoja.cell(row=fila, column=columna, value=valor)
                if columna >= 5:
                    celda.number_format = FORMATO_PESOS

        primera, ultima = fila_encabezado + 1, fila
        fila_total = ultima + 1
        hoja.cell(row=fila_total, column=1, value="Totales").font = Font(bold=True)
        letras = [get_column_letter(n) for n in range(4, len(ENCABEZADOS) + 1)]
        for columna, letra in enumerate(letras, start=4):
            celda = hoja.cell(row=fila_total, column=columna)
            celda.value = f"=SUM({letra}{primera}:{letra}{ultima})"
            celda.font = Font(bold=True)
            celda.fill = PatternFill("solid", fgColor=GRIS)
            if columna >= 5:
                celda.number_format = FORMATO_PESOS

        for columna, ancho in enumerate((11, 28, 14, 10, 20, 20, 22, 20, 20), start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def pdf(self, reporte: dict, ahora: datetime | None = None) -> bytes:
        """Documento PDF horizontal con la tabla de rentabilidad y los totales."""
        ahora = ahora or datetime.now()
        salida = BytesIO()
        documento = SimpleDocTemplate(
            salida,
            pagesize=landscape(A4),
            leftMargin=1.5 * cm,
            rightMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
            title="Reporte de rentabilidad - Pastoreando",
            author="Pastoreando",
            pageCompression=0,
        )
        estilos = getSampleStyleSheet()
        titulo = ParagraphStyle(
            "titulo", parent=estilos["Title"], textColor=colors.HexColor("#" + VERDE)
        )
        elementos = [
            Paragraph(_titulo(reporte), titulo),
            Paragraph(f"Generado el {ahora:%d/%m/%Y %H:%M}", estilos["Normal"]),
            Spacer(1, 0.5 * cm),
        ]

        filas = [list(ENCABEZADOS)]
        for c in reporte["contratos"]:
            filas.append(
                [
                    numero_contrato(c["id_contrato"]),
                    c["nombre_lote"],
                    c["estado"],
                    str(c["total_cabezas"]),
                    *(moneda(c[campo]) for campo in CAMPOS_VALOR),
                ]
            )
        totales = reporte["totales"]
        cabezas = sum(c["total_cabezas"] for c in reporte["contratos"])
        filas.append(
            ["Totales", "", "", str(cabezas)]
            + [moneda(totales[campo]) for campo in CAMPOS_VALOR]
        )

        tabla = Table(filas, repeatRows=1)
        tabla.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + VERDE)),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + GRIS)),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BDBDBD")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        elementos.append(tabla)
        if not reporte["contratos"]:
            elementos.append(
                Paragraph("No hay contratos para reportar.", estilos["Normal"])
            )
        documento.build(elementos)
        return salida.getvalue()
