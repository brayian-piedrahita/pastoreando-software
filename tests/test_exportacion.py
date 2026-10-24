"""Pruebas de la exportacion del reporte a Excel y PDF (GA8-AA1-EV02, RF06)."""

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.exportacion_servicio import (
    TIPO_EXCEL,
    TIPO_PDF,
    ExportacionServicio,
    moneda,
    nombre_archivo,
    numero_contrato,
)
from tests.conftest import (
    ARRENDATARIO,
    CONTADOR,
    OTRO_PROPIETARIO,
    PROPIETARIO,
    TECNICO,
)

RUTA_EXCEL = "/api/reportes/rentabilidad/excel"
RUTA_PDF = "/api/reportes/rentabilidad/pdf"


def _hoja(contenido: bytes):
    return load_workbook(BytesIO(contenido))["Rentabilidad"]


def test_excel_del_propietario_trae_sus_contratos(cliente, entrar):
    """CP-U01 | RF06 | Carlos descarga el Excel | encabezado, 2 contratos y totales"""
    r = cliente.get(RUTA_EXCEL, headers=entrar(PROPIETARIO))
    assert r.status_code == 200
    assert r.headers["content-type"] == TIPO_EXCEL
    assert 'attachment; filename="reporte_rentabilidad_' in (
        r.headers["content-disposition"]
    )
    hoja = _hoja(r.content)
    assert hoja["A4"].value == "Contrato"
    assert hoja["A5"].value == "0001"
    assert hoja["E5"].value == 13_260_000
    assert hoja["A7"].value == "Totales"
    assert hoja["E7"].value == "=SUM(E5:E6)"


def test_excel_del_contador_trae_los_cinco_contratos(cliente, entrar):
    """CP-U02 | RF06 | El contador descarga el Excel | cinco filas y fila de totales"""
    r = cliente.get(RUTA_EXCEL, headers=entrar(CONTADOR))
    hoja = _hoja(r.content)
    assert [hoja.cell(row=n, column=1).value for n in range(5, 11)] == [
        "0001",
        "0002",
        "0003",
        "0004",
        "0005",
        "Totales",
    ]


def test_pdf_es_un_documento_con_el_reporte(cliente, entrar):
    """CP-U03 | RF06 | Carlos descarga el PDF | documento PDF con titulo y contrato"""
    r = cliente.get(RUTA_PDF, headers=entrar(PROPIETARIO))
    assert r.status_code == 200
    assert r.headers["content-type"] == TIPO_PDF
    assert r.headers["content-disposition"].endswith('.pdf"')
    assert r.content.startswith(b"%PDF-")
    assert b"Reporte de rentabilidad" in r.content
    assert b"(0001)" in r.content


def test_pdf_del_otro_propietario_no_incluye_contratos_ajenos(cliente, entrar):
    """CP-U04 | RF06 | Marta descarga el PDF | solo sus contratos 3, 4 y 5"""
    r = cliente.get(RUTA_PDF, headers=entrar(OTRO_PROPIETARIO))
    assert r.status_code == 200
    assert b"(0003)" in r.content
    assert b"(0001)" not in r.content


@pytest.mark.parametrize("ruta", [RUTA_EXCEL, RUTA_PDF])
@pytest.mark.parametrize("correo", [ARRENDATARIO, TECNICO])
def test_exportar_con_rol_no_autorizado(cliente, entrar, ruta, correo):
    """CP-U05 | RF06 | Arrendatario o tecnico descargan | 403"""
    assert cliente.get(ruta, headers=entrar(correo)).status_code == 403


@pytest.mark.parametrize("ruta", [RUTA_EXCEL, RUTA_PDF])
def test_exportar_sin_sesion(cliente, ruta):
    """CP-U06 | RF06 | Descarga sin token | 401"""
    assert cliente.get(ruta).status_code == 401


def test_exportar_sin_contratos_genera_archivos_validos():
    """CP-U07 | RF06 | Reporte vacio | Excel sin filas y PDF con aviso"""
    vacio = {
        "margen_utilidad": 0.2,
        "contratos": [],
        "totales": dict.fromkeys(
            (
                "ingreso_total",
                "costo_estimado",
                "utilidad_proyectada",
                "recaudado",
                "saldo_por_cobrar",
            ),
            0.0,
        ),
    }
    servicio = ExportacionServicio()
    assert _hoja(servicio.excel(vacio))["A5"].value == "Totales"
    assert b"No hay contratos" in servicio.pdf(vacio)


def test_formatos_de_apoyo():
    """CP-U08 | RF06 | Formatos | moneda colombiana, numero de contrato y nombre"""
    assert moneda(1234567.5) == "$ 1.234.567,50"
    assert numero_contrato(7) == "0007"
    assert nombre_archivo("pdf", datetime(2026, 10, 24)) == (
        "reporte_rentabilidad_20261024.pdf"
    )
